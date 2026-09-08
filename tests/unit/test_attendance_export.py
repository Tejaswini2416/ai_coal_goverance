"""
Unit & Integration Tests for Statutory Attendance Spreadsheet Export (.xlsx and .csv).
"""
import io
import uuid
from datetime import datetime, timezone
import pytest
from httpx import AsyncClient, ASGITransport
import openpyxl

from app.main import app
from app.infrastructure.database.models import WorkerAttendanceModel, MineSiteModel, TenantModel
from app.services.auth_service import create_access_token


@pytest.mark.asyncio
async def test_attendance_export_csv_flow(mine_site, inspector_user):
    """Test exporting attendance records to CSV with DGMS statutory headers."""
    # Seed attendance records
    att1 = WorkerAttendanceModel(
        id=uuid.uuid4(),
        worker_id="W-101",
        worker_name="Rajesh Kumar Mandal",
        mine_site_id=mine_site.id,
        zone_type="UNDERGROUND",
        station_id="STATION-01",
        shift="MORNING",
        check_in_time=datetime.now(timezone.utc),
        biometric_verified=True,
        gas_level_exposure_ppm=12.5,
        status="ACTIVE_INSIDE",
    )
    await att1.insert()

    token = create_access_token(inspector_user)
    headers = {"Authorization": f"Bearer {token}"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.get(
            f"/api/v1/attendance/export?mine_site_id={mine_site.id}&format=csv",
            headers=headers,
        )
        assert resp.status_code == 200
        assert "text/csv" in resp.headers["content-type"]
        assert "attachment; filename=" in resp.headers["content-disposition"]
        
        content = resp.text
        assert "DIRECTORATE GENERAL OF MINES SAFETY" in content
        assert "Rajesh Kumar Mandal" in content
        assert "W-101" in content
        assert "UNDERGROUND" in content


@pytest.mark.asyncio
async def test_attendance_export_xlsx_flow(mine_site, inspector_user):
    """Test exporting attendance records to formatted Microsoft Excel (.xlsx)."""
    # Seed attendance records with gas hazard
    att1 = WorkerAttendanceModel(
        id=uuid.uuid4(),
        worker_id="W-102",
        worker_name="Suresh Goud",
        mine_site_id=mine_site.id,
        zone_type="UNDERGROUND",
        station_id="LEVEL3-SEAM4",
        shift="MORNING",
        check_in_time=datetime.now(timezone.utc),
        biometric_verified=True,
        gas_level_exposure_ppm=35.0,  # Elevated hazard
        status="ACTIVE_INSIDE",
    )
    await att1.insert()

    token = create_access_token(inspector_user)
    headers = {"Authorization": f"Bearer {token}"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.get(
            f"/api/v1/attendance/export?mine_site_id={mine_site.id}&format=xlsx",
            headers=headers,
        )
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
        assert ".xlsx" in resp.headers["content-disposition"]

        # Parse Excel workbook using openpyxl
        wb = openpyxl.load_workbook(filename=io.BytesIO(resp.content))
        ws = wb.active
        assert ws.title == "DGMS Attendance Register"
        assert "DIRECTORATE GENERAL OF MINES SAFETY" in str(ws["A1"].value)
        
        # Verify row 7 has Suresh Goud
        values = [str(ws.cell(row=7, column=col).value) for col in range(1, 12)]
        assert any("W-102" in v for v in values)
        assert any("Suresh Goud" in v for v in values)
        assert any("35.0 ppm" in v for v in values)
